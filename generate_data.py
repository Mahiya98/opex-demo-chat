"""Generate synthetic OPEX demo data for a fictional manufacturing company.

Creates CSV files, a combined JSON file, and a multi-tab Excel workbook that can
be uploaded to Google Drive and opened as a Google Sheet.
"""

import csv
import json
import math
import os
import random
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(42)

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "data")
os.makedirs(DATA, exist_ok=True)

COMPANY = "Vertex Demo Manufacturing"
DEPARTMENT = "Operational Excellence (OPEX)"
INDUSTRY = "Building materials (demonstration)"
START = date(2026, 1, 1)
END = date(2026, 10, 18)

PLANTS = ["Plant A", "Plant B", "Plant C"]
LINES = ["Line 1", "Line 2"]
SHIFTS = ["A", "B"]
ZONES = ["Receiving", "Line 1", "Line 2", "Packing", "Warehouse"]
PRODUCTS = ["Product A", "Product B"]
AUDITORS = ["R. Das", "M. Iqbal", "S. Roy", "T. Ahmed", "K. Nair"]
OWNERS = ["R. Das", "M. Iqbal", "S. Roy", "T. Ahmed", "K. Nair", "P. Sen"]
DEPARTMENTS = ["Production", "Maintenance", "Quality", "Packing", "Utilities", "Safety"]
KAIZEN_CATEGORIES = ["Safety", "Quality", "Cost", "Delivery", "Morale", "Energy"]
BREAKDOWN_CATEGORIES = ["Mechanical", "Electrical", "Instrumentation", "Utility", "Operational"]
MACHINES = [
    ("Plant A", "Mill-01"), ("Plant A", "Mill-02"), ("Plant A", "Kiln-01"),
    ("Plant A", "Packer-01"), ("Plant A", "Packer-02"), ("Plant A", "Conveyor-02"),
    ("Plant B", "Mill-01"), ("Plant B", "Kiln-01"), ("Plant B", "Packer-01"),
    ("Plant B", "Compressor-01"), ("Plant B", "Conveyor-03"),
    ("Plant C", "Mill-01"), ("Plant C", "Kiln-02"), ("Plant C", "Packer-01"),
    ("Plant C", "Compressor-02"), ("Plant C", "Conveyor-01"),
]

BASE_PLAN = {
    ("Plant A", "Line 1"): 420, ("Plant A", "Line 2"): 380,
    ("Plant B", "Line 1"): 360, ("Plant B", "Line 2"): 330,
    ("Plant C", "Line 1"): 310, ("Plant C", "Line 2"): 290,
}
COST_PER_TONNE = {"Plant A": 56, "Plant B": 59, "Plant C": 62}

SECTION_TESTS = {
    "Milling": [
        ("Fineness", "m2/kg", 300, 380, 340),
        ("Residue", "%", 0, 8, 3.8),
    ],
    "Kiln": [
        ("Moisture", "%", 0, 1.5, 0.75),
        ("Setting Time", "min", 90, 180, 128),
    ],
    "Packing": [
        ("Packing Weight", "kg", 49.5, 50.5, 50.0),
    ],
    "Quality Lab": [
        ("Compressive Strength", "MPa", 42, 60, 50.5),
    ],
}
SECTION_MACHINES = {
    "Milling": ["Mill-01", "Mill-02"],
    "Kiln": ["Kiln-01", "Kiln-02"],
    "Packing": ["Packer-01", "Packer-02"],
    "Quality Lab": ["Lab Press"],
}

KAIZEN_TITLES = [
    "Reduce bag burst in packing area",
    "Optimize mill feed rate",
    "Eliminate water leak in utility line",
    "Shorten kiln warm-up time",
    "Improve pallet stacking method",
    "Reduce compressed-air leakage",
    "Standardize shift handover checklist",
    "Reduce material spillage at conveyor transfer",
    "Improve sample labelling in quality lab",
    "Reduce changeover time on line 2",
    "Install visual level indicator on silo",
    "Optimize lighting schedule in warehouse",
]
BREAKDOWN_CAUSES = [
    "Bearing failure", "Belt slip", "Motor overload", "Sensor fault",
    "Hydraulic leak", "Blocked chute", "Power trip", "Coupling wear",
    "Cooling fan failure", "Control card fault",
]


def clamp(value, low, high):
    return max(low, min(high, value))


def season(d):
    return 0.95 + 0.08 * math.sin((d.month - 1) / 12 * 2 * math.pi)


def weekdays():
    d = START
    while d <= END:
        if d.weekday() != 6:
            yield d
        d += timedelta(days=1)


def write_csv(name, header, rows):
    path = os.path.join(DATA, name + ".csv")
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)
    return path


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for idx, column in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[get_column_letter(idx)].width = min(42, max(11, width + 2))


# ---------------------------------------------------------------- production
production_rows = []
for d in weekdays():
    for plant in PLANTS:
        for line in LINES:
            for shift in SHIFTS:
                planned = BASE_PLAN[(plant, line)] * season(d)
                if shift == "B":
                    planned *= 0.92
                planned = int(round(planned))
                factor = clamp(random.gauss(0.965, 0.055), 0.72, 1.08)
                actual = int(round(planned * factor))
                downtime = round(clamp(random.gauss(0.7, 0.6), 0, 4.5), 2)
                runtime = round(8 - downtime, 2)
                reject_rate = clamp(random.gauss(0.022, 0.012), 0.005, 0.08)
                good = int(round(actual * (1 - reject_rate)))
                reject = max(0, actual - good)
                production_rows.append([
                    d.isoformat(), plant, line, shift,
                    random.choice(PRODUCTS), planned, actual, good, reject,
                    runtime, downtime,
                ])

# --------------------------------------------------------------------- 5S
five_s_rows = []
plant_five_s = {"Plant A": 3.9, "Plant B": 3.6, "Plant C": 3.25}
zone_bias = {
    "Receiving": -0.15, "Line 1": 0.10, "Line 2": 0.05,
    "Packing": -0.10, "Warehouse": -0.25,
}
d = START
week_no = 0
while d <= END:
    if d.weekday() == 0:
        week_no += 1
        for plant in PLANTS:
            for zone in ZONES:
                trend = week_no * 0.012
                base = plant_five_s[plant] + zone_bias[zone] + trend

                def score():
                    return round(clamp(base + random.gauss(0, 0.28), 1, 5), 2)

                sort_removed = score()
                sort_redtag = score()
                order_location = score()
                order_labels = score()
                shine_clean = score()
                shine_assigned = score()
                standardize = score()
                sustain = score()
                total = (
                    (sort_removed + sort_redtag) / 2
                    + (order_location + order_labels) / 2
                    + (shine_clean + shine_assigned) / 2
                    + standardize
                    + sustain
                )
                score_pct = round(total / 25 * 100, 2)
                remarks = ""
                if score_pct < 60:
                    remarks = random.choice([
                        "Cleaning pending after maintenance",
                        "Unlabelled material near line",
                        "Tools not returned to shadow board",
                    ])
                five_s_rows.append([
                    f"5S-{len(five_s_rows) + 1:04d}", d.isoformat(), plant, zone,
                    random.choice(AUDITORS), sort_removed, sort_redtag,
                    order_location, order_labels, shine_clean, shine_assigned,
                    standardize, sustain, score_pct, remarks,
                ])
    d += timedelta(days=1)

# --------------------------------------------------------------------- QCP
qcp_rows = []
machine_bias = {
    "Mill-01": 4, "Mill-02": -3, "Kiln-01": 0, "Kiln-02": -1,
    "Packer-01": 0.08, "Packer-02": -0.07, "Lab Press": 0,
}
for d in weekdays():
    for plant in PLANTS:
        sections = random.sample(list(SECTION_TESTS.keys()), 2)
        for section in sections:
            test_name, unit, low, high, base_value = random.choice(SECTION_TESTS[section])
            machine = random.choice(SECTION_MACHINES[section])
            shift = random.choice(SHIFTS)
            spread = (high - low) * 0.22
            value = base_value + machine_bias.get(machine, 0) + random.gauss(0, spread)
            if test_name == "Packing Weight":
                value = round(value, 2)
            elif test_name in ("Moisture", "Residue"):
                value = round(value, 2)
            else:
                value = round(value, 1)
            status = "Pass" if low <= value <= high else "Fail"
            remarks = ""
            if status == "Fail":
                remarks = random.choice([
                    "Out of spec, recheck requested",
                    "Adjustment made and sample retained",
                    "Shift informed, parameter corrected",
                ])
            qcp_rows.append([
                f"QCP-{len(qcp_rows) + 1:05d}", d.isoformat(), plant, section,
                machine, random.choice(PRODUCTS), shift, test_name, unit,
                low, high, value, status, remarks,
            ])

# ------------------------------------------------------------------ Kaizen
kaizen_rows = []
for month in range(1, 11):
    for plant in PLANTS:
        for _ in range(random.randint(6, 9)):
            day = random.randint(1, 28)
            if month == 10:
                day = min(day, 18)
            created = date(2026, month, day)
            status = random.choices(
                ["Completed", "In Progress", "Planned"], weights=[62, 28, 10]
            )[0]
            completion = ""
            if status == "Completed":
                completion = (created + timedelta(days=random.randint(5, 30))).isoformat()
            kaizen_rows.append([
                f"KZ-{len(kaizen_rows) + 1:04d}", created.isoformat(), plant,
                random.choice(DEPARTMENTS), random.choice(KAIZEN_TITLES),
                random.choice(KAIZEN_CATEGORIES), status, random.choice(OWNERS),
                round(random.uniform(500, 25000), 2), completion,
            ])
kaizen_rows.sort(key=lambda row: row[1])

# --------------------------------------------------------------- Breakdown
breakdown_rows = []
for _ in range(320):
    plant, machine = random.choice(MACHINES)
    day = START + timedelta(days=random.randint(0, (END - START).days))
    start_hour = random.randint(0, 22)
    start_time = datetime(day.year, day.month, day.day, start_hour, random.choice([0, 15, 30, 45]))
    duration = round(clamp(random.expovariate(0.35), 0.5, 12), 2)
    end_time = start_time + timedelta(hours=duration)
    rate = 22 if "Mill" in machine else 14 if "Kiln" in machine else 8
    loss = int(round(duration * rate * random.uniform(0.7, 1.2)))
    breakdown_rows.append([
        f"BD-{len(breakdown_rows) + 1:04d}", day.isoformat(), plant, machine,
        random.choice(BREAKDOWN_CATEGORIES), random.choice(BREAKDOWN_CAUSES),
        start_time.strftime("%H:%M"), end_time.strftime("%H:%M"), duration,
        loss, round(random.uniform(150, 8000), 2),
        random.choices(["Closed", "Open"], weights=[92, 8])[0], random.choice(OWNERS),
    ])
breakdown_rows.sort(key=lambda row: row[1])

# ----------------------------------------------------------- Plan variance
monthly = defaultdict(lambda: {"planned": 0, "actual": 0})
for row in production_rows:
    key = (row[0][:7], row[1])
    monthly[key]["planned"] += row[5]
    monthly[key]["actual"] += row[6]

plan_rows = []
for (month, plant), values in sorted(monthly.items()):
    planned = values["planned"]
    actual = values["actual"]
    variance = actual - planned
    variance_pct = round(variance / planned * 100, 2) if planned else 0
    planned_cost = round(planned * COST_PER_TONNE[plant], 2)
    actual_cost = round(actual * COST_PER_TONNE[plant] * random.uniform(0.97, 1.05), 2)
    cost_var = round((actual_cost - planned_cost) / planned_cost * 100, 2) if planned_cost else 0
    plan_rows.append([
        month, plant, planned, actual, variance, variance_pct,
        planned_cost, actual_cost, cost_var,
    ])

# -------------------------------------------------------------------- CSVs
write_csv(
    "production",
    ["date", "plant", "line", "shift", "product", "planned_tonnes", "actual_tonnes",
     "good_tonnes", "reject_tonnes", "runtime_hours", "downtime_hours"],
    production_rows,
)
write_csv(
    "5s_audit",
    ["audit_id", "date", "plant", "zone", "auditor", "sort_removed", "sort_redtag",
     "order_location", "order_labels", "shine_clean", "shine_assigned",
     "standardize", "sustain", "score_pct", "remarks"],
    five_s_rows,
)
write_csv(
    "qcp_audit",
    ["audit_id", "date", "plant", "section", "machine", "product", "shift",
     "test_name", "unit", "min_value", "max_value", "actual_value", "status", "remarks"],
    qcp_rows,
)
write_csv(
    "kaizen",
    ["kaizen_id", "date", "plant", "department", "title", "category", "status",
     "owner", "savings_usd", "completion_date"],
    kaizen_rows,
)
write_csv(
    "breakdown",
    ["breakdown_id", "date", "plant", "machine", "category", "cause", "start_time",
     "end_time", "downtime_hours", "production_loss_tonnes", "maintenance_cost_usd",
     "status", "owner"],
    breakdown_rows,
)
write_csv(
    "plan_variance",
    ["month", "plant", "planned_tonnes", "actual_tonnes", "variance_tonnes",
     "variance_pct", "planned_cost_usd", "actual_cost_usd", "cost_variance_pct"],
    plan_rows,
)

# -------------------------------------------------------------------- JSON
HEADERS = {
    "production": ["date", "plant", "line", "shift", "product", "planned_tonnes",
                   "actual_tonnes", "good_tonnes", "reject_tonnes", "runtime_hours",
                   "downtime_hours"],
    "5s_audit": ["audit_id", "date", "plant", "zone", "auditor", "sort_removed",
                 "sort_redtag", "order_location", "order_labels", "shine_clean",
                 "shine_assigned", "standardize", "sustain", "score_pct", "remarks"],
    "qcp_audit": ["audit_id", "date", "plant", "section", "machine", "product", "shift",
                  "test_name", "unit", "min_value", "max_value", "actual_value",
                  "status", "remarks"],
    "kaizen": ["kaizen_id", "date", "plant", "department", "title", "category",
               "status", "owner", "savings_usd", "completion_date"],
    "breakdown": ["breakdown_id", "date", "plant", "machine", "category", "cause",
                  "start_time", "end_time", "downtime_hours", "production_loss_tonnes",
                  "maintenance_cost_usd", "status", "owner"],
    "plan_variance": ["month", "plant", "planned_tonnes", "actual_tonnes",
                      "variance_tonnes", "variance_pct", "planned_cost_usd",
                      "actual_cost_usd", "cost_variance_pct"],
}


def as_objects(key, rows):
    return [dict(zip(HEADERS[key], row)) for row in rows]


payload = {
    "meta": {
        "company": COMPANY,
        "department": DEPARTMENT,
        "industry": INDUSTRY,
        "note": "Synthetic demo data. All names, figures and events are fictional.",
        "from": START.isoformat(),
        "to": END.isoformat(),
        "generated": datetime.now().isoformat(timespec="seconds"),
    },
    "production": as_objects("production", production_rows),
    "5s_audit": as_objects("5s_audit", five_s_rows),
    "qcp_audit": as_objects("qcp_audit", qcp_rows),
    "kaizen": as_objects("kaizen", kaizen_rows),
    "breakdown": as_objects("breakdown", breakdown_rows),
    "plan_variance": as_objects("plan_variance", plan_rows),
}
with open(os.path.join(DATA, "demo-data.json"), "w", encoding="utf-8") as handle:
    json.dump(payload, handle, separators=(",", ":"))

# ------------------------------------------------------------------- Excel
wb = Workbook()
about = wb.active
about.title = "About"
about.append(["Field", "Value"])
about.append(["Company", COMPANY])
about.append(["Department", DEPARTMENT])
about.append(["Industry", INDUSTRY])
about.append(["Period", f"{START.isoformat()} to {END.isoformat()}"])
about.append(["Data", "Synthetic demonstration data - all values are fictional"])
about.append(["Rows", f"Production {len(production_rows)}, 5S {len(five_s_rows)}, QCP {len(qcp_rows)}, Kaizen {len(kaizen_rows)}, Breakdown {len(breakdown_rows)}"])
style_sheet(about)

sheets = [
    ("Production", ["date", "plant", "line", "shift", "product", "planned_tonnes",
                    "actual_tonnes", "good_tonnes", "reject_tonnes", "runtime_hours",
                    "downtime_hours"], production_rows),
    ("5S_Audit", ["audit_id", "date", "plant", "zone", "auditor", "sort_removed",
                  "sort_redtag", "order_location", "order_labels", "shine_clean",
                  "shine_assigned", "standardize", "sustain", "score_pct", "remarks"],
     five_s_rows),
    ("QCP_Audit", ["audit_id", "date", "plant", "section", "machine", "product", "shift",
                   "test_name", "unit", "min_value", "max_value", "actual_value",
                   "status", "remarks"], qcp_rows),
    ("Kaizen", ["kaizen_id", "date", "plant", "department", "title", "category",
                "status", "owner", "savings_usd", "completion_date"], kaizen_rows),
    ("Breakdown", ["breakdown_id", "date", "plant", "machine", "category", "cause",
                   "start_time", "end_time", "downtime_hours", "production_loss_tonnes",
                   "maintenance_cost_usd", "status", "owner"], breakdown_rows),
    ("Plan_Variance", ["month", "plant", "planned_tonnes", "actual_tonnes",
                       "variance_tonnes", "variance_pct", "planned_cost_usd",
                       "actual_cost_usd", "cost_variance_pct"], plan_rows),
]
for title, header, rows in sheets:
    ws = wb.create_sheet(title)
    ws.append(header)
    for row in rows:
        ws.append(row)
    style_sheet(ws)

wb.save(os.path.join(BASE, "Demo_OPEX_Data.xlsx"))

print("Generated:")
print(f"  production : {len(production_rows)} rows")
print(f"  5S audit   : {len(five_s_rows)} rows")
print(f"  QCP audit  : {len(qcp_rows)} rows")
print(f"  kaizen     : {len(kaizen_rows)} rows")
print(f"  breakdown  : {len(breakdown_rows)} rows")
print(f"  plan var   : {len(plan_rows)} rows")
print(f"  workbook   : {os.path.join(BASE, 'Demo_OPEX_Data.xlsx')}")
